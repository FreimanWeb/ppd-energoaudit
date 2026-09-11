"""Причины простоев нагнетательных скважин.

Промысловая сводка «Агент закачки»: по строке на скважину, с причиной
простоя и датой его начала. Кладётся в ``data/downtime/<объект>.xlsx`` как
есть, без переработки — файл приходит из отчётности и обновляется целиком.

Шапка таблицы на второй строке; ниже у скважины могут идти строки-продолжения
по пластам, где номер скважины пуст, — они к простою отношения не имеют.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any


DOWNTIME_DIRNAME = "downtime"
HEADER_ROW = 2

WELL_COLUMN = "Скваж"
REASON_COLUMN = "Причина_простоя"
SINCE_COLUMN = "Дата_непр_пр"
RUNTIME_COLUMN = "T_раб"


class DowntimeError(ValueError):
    """Сводка простоев не читается."""


@dataclass(frozen=True)
class Downtime:
    """Причина простоя одной скважины."""

    well: str
    reason: str
    since: datetime | None = None
    runtime_hours: float | None = None


def normalize_well(name: str) -> str:
    """Номер скважины для сопоставления с моделью.

    В сводке встречается «3592Д/2», в выгрузке CRM — «3592Д»: хвост после
    косой черты означает ствол, а не отдельную скважину.
    """
    return str(name or "").strip().split("/")[0].strip()


def _as_datetime(value: Any) -> datetime | None:
    """Момент начала простоя; в сводке он указан с точностью до минуты."""
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
        "%d.%m.%Y %H:%M:%S",
        "%d.%m.%Y %H:%M",
        "%d.%m.%Y",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _as_float(value: Any) -> float | None:
    text = str(value or "").strip().replace(",", ".")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return None if number != number else number


def read_downtime(path: Path) -> list[Downtime]:
    """Скважины с указанной причиной простоя; остальные строки пропускаются."""
    import openpyxl

    if not path.exists():
        return []
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if len(rows) < HEADER_ROW:
        raise DowntimeError(f"в {path.name} нет шапки таблицы")
    header = [str(cell).strip() if cell is not None else "" for cell in rows[HEADER_ROW - 1]]
    if WELL_COLUMN not in header or REASON_COLUMN not in header:
        raise DowntimeError(
            f"в {path.name} нет столбцов «{WELL_COLUMN}» и «{REASON_COLUMN}»"
        )

    index = {name: header.index(name) for name in header if name}

    def cell(row: tuple, name: str) -> Any:
        position = index.get(name)
        return row[position] if position is not None and position < len(row) else None

    out = []
    for row in rows[HEADER_ROW:]:
        well = normalize_well(cell(row, WELL_COLUMN))
        reason = str(cell(row, REASON_COLUMN) or "").strip()
        if not well or not reason:
            continue
        out.append(
            Downtime(
                well=well,
                reason=reason,
                since=_as_datetime(cell(row, SINCE_COLUMN)),
                runtime_hours=_as_float(cell(row, RUNTIME_COLUMN)),
            )
        )
    return out


def by_well(records: list[Downtime]) -> dict[str, Downtime]:
    """Последняя запись на скважину: в сводке скважина встречается один раз."""
    return {record.well: record for record in records}


def started_in(records: list[Downtime], first: date, last: date) -> list[Downtime]:
    """Простои, начавшиеся внутри периода, — их и есть смысл класть на график.

    Давние простои (консервация с 2022 года) внутрь окна не попадают: на
    кривой они ничего не объясняют, для них есть столбец в таблице.
    """
    return sorted(
        (
            record
            for record in records
            if record.since is not None and first <= record.since.date() <= last
        ),
        key=lambda record: (record.since, record.well),
    )
